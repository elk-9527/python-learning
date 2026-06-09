import os
import math
import numpy as np
import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


# =========================
# 基础设置
# =========================

OUT_DIR = "ppt_assets"
os.makedirs(OUT_DIR, exist_ok=True)

PPT_NAME = "数值积分算法设计与实现_汇报PPT.pptx"

CH_FONT = "Microsoft YaHei"

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["mathtext.fontset"] = "dejavusans"

BG = RGBColor(248, 250, 252)
DARK = RGBColor(30, 41, 59)
SUB = RGBColor(71, 85, 105)
BLUE = RGBColor(37, 99, 235)
ORANGE = RGBColor(249, 115, 22)
GREEN = RGBColor(22, 163, 74)
RED = RGBColor(220, 38, 38)
PURPLE = RGBColor(124, 58, 237)
GRAY = RGBColor(226, 232, 240)
WHITE = RGBColor(255, 255, 255)


# =========================
# 数值积分计算
# =========================

def v(t):
    return 4.0 / (1.0 + t * t)


EXACT = math.pi


def trapezoid_single(f, a, b):
    return (b - a) * (f(a) + f(b)) / 2.0


def simpson_single(f, a, b):
    m = (a + b) / 2.0
    return (b - a) * (f(a) + 4.0 * f(m) + f(b)) / 6.0


def composite_trapezoid(f, a, b, n):
    h = (b - a) / n
    total = 0.5 * (f(a) + f(b))
    for i in range(1, n):
        total += f(a + i * h)
    return h * total


def composite_simpson(f, a, b, n):
    if n % 2 != 0:
        raise ValueError("复化辛普森公式要求 n 必须为偶数。")
    h = (b - a) / n
    odd_sum = 0.0
    even_sum = 0.0
    for i in range(1, n):
        x = a + i * h
        if i % 2 == 1:
            odd_sum += f(x)
        else:
            even_sum += f(x)
    return h / 3.0 * (f(a) + f(b) + 4.0 * odd_sum + 2.0 * even_sum)


a, b = 0.0, 1.0
ns = [2, 4, 8, 16, 32, 64]
hs = [1 / n for n in ns]

T_single = trapezoid_single(v, a, b)
S_single = simpson_single(v, a, b)

err_T_single = abs(T_single - EXACT)
err_S_single = abs(S_single - EXACT)

trap_values = [composite_trapezoid(v, a, b, n) for n in ns]
simp_values = [composite_simpson(v, a, b, n) for n in ns]

trap_errors = [abs(x - EXACT) for x in trap_values]
simp_errors = [abs(x - EXACT) for x in simp_values]

costs = [n + 1 for n in ns]


# =========================
# 绘图函数
# =========================

def save_velocity_area():
    t = np.linspace(0, 1, 500)
    y = v(t)

    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=180)
    ax.plot(t, y, color="#2563eb", lw=3, label=r"$v(t)=\frac{4}{1+t^2}$")
    ax.fill_between(t, 0, y, color="#93c5fd", alpha=0.55)
    ax.scatter([0, 0.5, 1], [v(0), v(0.5), v(1)], color="#f97316", s=60, zorder=5)

    ax.text(0.48, 1.25, r"$S=\int_0^1 \frac{4}{1+t^2}\,dt=\pi$",
            fontsize=17, color="#1e293b")

    ax.set_title("速度曲线与位移面积", fontsize=16)
    ax.set_xlabel("时间 t / s")
    ax.set_ylabel("速度 v(t) / m·s$^{-1}$")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 4.4)
    ax.grid(True, linestyle="--", alpha=0.35)
    ax.legend()

    path = os.path.join(OUT_DIR, "velocity_area.png")
    plt.tight_layout()
    plt.savefig(path, bbox_inches="tight")
    plt.close()
    return path


def save_method_sketch():
    xs = np.linspace(0, 1, 400)
    ys = v(xs)

    x_nodes = np.array([0, 1])
    y_nodes = v(x_nodes)
    line_y = np.interp(xs, x_nodes, y_nodes)

    x_simp = np.array([0, 0.5, 1])
    y_simp = v(x_simp)
    coef = np.polyfit(x_simp, y_simp, 2)
    poly_y = np.polyval(coef, xs)

    fig, axes = plt.subplots(1, 2, figsize=(10, 4.2), dpi=180)

    ax = axes[0]
    ax.plot(xs, ys, lw=3, color="#2563eb", label="真实曲线")
    ax.plot(xs, line_y, lw=2.5, color="#f97316", label="直线近似")
    ax.fill_between(xs, 0, line_y, color="#fdba74", alpha=0.45)
    ax.scatter(x_nodes, y_nodes, color="#f97316", s=55)
    ax.set_title("梯形公式：直线近似")
    ax.set_xlabel("t")
    ax.set_ylabel("v(t)")
    ax.grid(True, linestyle="--", alpha=0.3)
    ax.legend(fontsize=9)

    ax = axes[1]
    ax.plot(xs, ys, lw=3, color="#2563eb", label="真实曲线")
    ax.plot(xs, poly_y, lw=2.5, color="#16a34a", label="抛物线近似")
    ax.fill_between(xs, 0, poly_y, color="#86efac", alpha=0.45)
    ax.scatter(x_simp, y_simp, color="#16a34a", s=55)
    ax.set_title("辛普森公式：抛物线近似")
    ax.set_xlabel("t")
    ax.set_ylabel("v(t)")
    ax.grid(True, linestyle="--", alpha=0.3)
    ax.legend(fontsize=9)

    path = os.path.join(OUT_DIR, "method_sketch.png")
    plt.tight_layout()
    plt.savefig(path, bbox_inches="tight")
    plt.close()
    return path


def save_single_error_bar():
    labels = ["普通梯形", "普通辛普森", "复化梯形\nn=16", "复化辛普森\nn=4"]
    values = [
        err_T_single,
        err_S_single,
        trap_errors[ns.index(16)],
        simp_errors[ns.index(4)]
    ]
    colors = ["#ef4444", "#22c55e", "#2563eb", "#f97316"]

    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=180)
    bars = ax.bar(labels, values, color=colors, alpha=0.9)
    ax.set_yscale("log")
    ax.set_ylabel("绝对误差")
    ax.set_title("代表性算法误差对比")
    ax.grid(True, axis="y", linestyle="--", alpha=0.35)

    for bar, val in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, val * 1.25,
                f"{val:.2e}", ha="center", fontsize=10)

    path = os.path.join(OUT_DIR, "single_error_bar.png")
    plt.tight_layout()
    plt.savefig(path, bbox_inches="tight")
    plt.close()
    return path


def save_error_step():
    fig, ax = plt.subplots(figsize=(8.5, 5), dpi=180)

    ax.loglog(hs, trap_errors, "o-", lw=2.5, ms=7, color="#2563eb", label="复化梯形")
    ax.loglog(hs, simp_errors, "s-", lw=2.5, ms=7, color="#f97316", label="复化辛普森")

    ax.hlines(err_T_single, min(hs), max(hs),
              colors="#ef4444", linestyles="--", lw=2, label="普通梯形误差")
    ax.hlines(err_S_single, min(hs), max(hs),
              colors="#22c55e", linestyles="--", lw=2, label="普通辛普森误差")

    ref_trap = [trap_errors[0] * (h / hs[0]) ** 2 for h in hs]
    ref_simp4 = [simp_errors[0] * (h / hs[0]) ** 4 for h in hs]
    ref_simp6 = [simp_errors[0] * (h / hs[0]) ** 6 for h in hs]

    ax.loglog(hs, ref_trap, "k--", alpha=0.45, lw=1.8, label=r"$O(h^2)$")
    ax.loglog(hs, ref_simp4, ":", color="#64748b", alpha=0.8, lw=2, label=r"$O(h^4)$")
    ax.loglog(hs, ref_simp6, "--", color="#7c3aed", alpha=0.6, lw=2, label=r"本题 $O(h^6)$")

    ax.set_title("误差—步长图", fontsize=16)
    ax.set_xlabel("步长 h = 1/n")
    ax.set_ylabel("绝对误差")
    ax.grid(True, which="both", linestyle="--", alpha=0.35)
    ax.legend(fontsize=9)
    ax.invert_xaxis()

    path = os.path.join(OUT_DIR, "error_step.png")
    plt.tight_layout()
    plt.savefig(path, bbox_inches="tight")
    plt.close()
    return path


def save_error_cost():
    fig, ax = plt.subplots(figsize=(8.5, 5), dpi=180)

    ax.loglog(2, err_T_single, "o", ms=10, color="#ef4444", label="普通梯形")
    ax.loglog(3, err_S_single, "o", ms=10, color="#22c55e", label="普通辛普森")
    ax.loglog(costs, trap_errors, "o-", lw=2.5, ms=7, color="#2563eb", label="复化梯形")
    ax.loglog(costs, simp_errors, "s-", lw=2.5, ms=7, color="#f97316", label="复化辛普森")

    for c, e, n in zip(costs, trap_errors, ns):
        ax.annotate(f"n={n}", (c, e), xytext=(6, 5),
                    textcoords="offset points", fontsize=9)

    ax.set_title("误差—计算量对比图", fontsize=16)
    ax.set_xlabel("函数计算次数 / 节点数")
    ax.set_ylabel("绝对误差")
    ax.grid(True, which="both", linestyle="--", alpha=0.35)
    ax.legend(fontsize=9)

    path = os.path.join(OUT_DIR, "error_cost.png")
    plt.tight_layout()
    plt.savefig(path, bbox_inches="tight")
    plt.close()
    return path


def save_convergence_ratio():
    labels = [f"{ns[i]}→{ns[i+1]}" for i in range(len(ns) - 1)]
    trap_ratio = [trap_errors[i] / trap_errors[i + 1] for i in range(len(ns) - 1)]
    simp_ratio = [simp_errors[i] / simp_errors[i + 1] for i in range(len(ns) - 1)]

    x = np.arange(len(labels))
    w = 0.36

    fig, ax = plt.subplots(figsize=(8.5, 4.8), dpi=180)
    ax.bar(x - w / 2, trap_ratio, width=w, color="#2563eb", label="复化梯形")
    ax.bar(x + w / 2, simp_ratio, width=w, color="#f97316", label="复化辛普森")

    ax.axhline(4, color="#2563eb", linestyle="--", alpha=0.55, label="二阶理论≈4")
    ax.axhline(64, color="#7c3aed", linestyle="--", alpha=0.55, label="本题六阶≈64")

    ax.set_yscale("log")
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylabel("误差缩小倍数")
    ax.set_title("n 翻倍时误差缩小倍数")
    ax.grid(True, axis="y", linestyle="--", alpha=0.35)
    ax.legend(fontsize=9)

    path = os.path.join(OUT_DIR, "convergence_ratio.png")
    plt.tight_layout()
    plt.savefig(path, bbox_inches="tight")
    plt.close()
    return path


def save_engineering_threshold():
    fig, ax = plt.subplots(figsize=(8.5, 5), dpi=180)

    ax.semilogy(ns, trap_errors, "o-", lw=2.5, ms=7, color="#2563eb", label="复化梯形")
    ax.semilogy(ns, simp_errors, "s-", lw=2.5, ms=7, color="#f97316", label="复化辛普森")

    ax.axhline(1e-3, color="#ef4444", linestyle="--", lw=2, label="工程要求 $10^{-3}$ m")
    ax.axhline(1e-4, color="#7c3aed", linestyle="--", lw=2, label="提高要求 $10^{-4}$ m")

    ax.set_xticks(ns)
    ax.set_xlabel("等分数 n")
    ax.set_ylabel("绝对误差 / m")
    ax.set_title("工程精度达标判断")
    ax.grid(True, which="both", linestyle="--", alpha=0.35)
    ax.legend(fontsize=9)

    path = os.path.join(OUT_DIR, "engineering_threshold.png")
    plt.tight_layout()
    plt.savefig(path, bbox_inches="tight")
    plt.close()
    return path


velocity_img = save_velocity_area()
method_img = save_method_sketch()
bar_img = save_single_error_bar()
error_step_img = save_error_step()
error_cost_img = save_error_cost()
ratio_img = save_convergence_ratio()
eng_img = save_engineering_threshold()


# =========================
# PPT 工具函数
# =========================

def add_background(slide):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0, 0,
        Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.color.rgb = BG

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0, 0,
        Inches(0.12), Inches(7.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.color.rgb = BLUE


def set_text_style(run, size=18, color=DARK, bold=False):
    run.font.name = CH_FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.18), Inches(12.2), Inches(0.55))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_text_style(run, 26, DARK, True)

    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.58), Inches(0.75), Inches(12), Inches(0.35))
        tf2 = box2.text_frame
        tf2.clear()
        p2 = tf2.paragraphs[0]
        run2 = p2.add_run()
        run2.text = subtitle
        set_text_style(run2, 13, SUB, False)


def add_footer(slide, page):
    box = slide.shapes.add_textbox(Inches(11.7), Inches(7.05), Inches(1.2), Inches(0.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{page:02d}"
    set_text_style(run, 10, SUB, False)


def add_card(slide, left, top, width, height, title, value, color=BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(203, 213, 225)

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top),
        Inches(0.08), Inches(height)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.color.rgb = color

    box1 = slide.shapes.add_textbox(Inches(left + 0.18), Inches(top + 0.14), Inches(width - 0.3), Inches(0.25))
    tf1 = box1.text_frame
    tf1.clear()
    p1 = tf1.paragraphs[0]
    r1 = p1.add_run()
    r1.text = title
    set_text_style(r1, 11, SUB, False)

    box2 = slide.shapes.add_textbox(Inches(left + 0.18), Inches(top + 0.45), Inches(width - 0.3), Inches(height - 0.5))
    tf2 = box2.text_frame
    tf2.clear()
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    r2.text = value
    set_text_style(r2, 18, DARK, True)


def add_text_box(slide, left, top, width, height, text, size=16, color=DARK, bold=False, align="left"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]

    if align == "center":
        p.alignment = PP_ALIGN.CENTER
    elif align == "right":
        p.alignment = PP_ALIGN.RIGHT
    else:
        p.alignment = PP_ALIGN.LEFT

    r = p.add_run()
    r.text = text
    set_text_style(r, size, color, bold)
    return box


def add_table(slide, data, left, top, width, height, font_size=11,
              header_color=BLUE, header_font=WHITE):
    rows = len(data)
    cols = len(data[0])
    table_shape = slide.shapes.add_table(
        rows, cols,
        Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    table = table_shape.table

    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = str(data[r][c])
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)

            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_color
                font_color = header_font
                bold = True
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if r % 2 == 1 else RGBColor(241, 245, 249)
                font_color = DARK
                bold = False

            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.name = CH_FONT
                    run.font.size = Pt(font_size)
                    run.font.bold = bold
                    run.font.color.rgb = font_color

    return table_shape


def add_picture(slide, path, left, top, width=None, height=None):
    if width and height:
        return slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(width), Inches(height))
    elif width:
        return slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(width))
    elif height:
        return slide.shapes.add_picture(path, Inches(left), Inches(top), height=Inches(height))
    else:
        return slide.shapes.add_picture(path, Inches(left), Inches(top))


def add_flow_box(slide, left, top, width, height, text, color):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color

    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_text_style(r, 15, WHITE, True)
    return shape


def add_arrow(slide, left, top, width, height):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(148, 163, 184)
    shape.line.color.rgb = RGBColor(148, 163, 184)
    return shape


# =========================
# 创建 PPT
# =========================

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


# Slide 1 封面
slide = prs.slides.add_slide(blank)
add_background(slide)

add_text_box(slide, 0.8, 0.65, 11.8, 0.7,
             "数值积分算法设计与实现", 36, DARK, True, "center")
add_text_box(slide, 1.1, 1.35, 11.2, 0.45,
             "梯形公式 · 辛普森公式 · 复化求积 · 工程误差分析", 18, SUB, False, "center")

add_picture(slide, velocity_img, 1.0, 2.05, width=6.3)

add_card(slide, 7.8, 2.15, 4.4, 0.9, "工程问题", "滑块位移积分", BLUE)
add_card(slide, 7.8, 3.25, 4.4, 0.9, "精确值", "S = π ≈ 3.1415926536 m", GREEN)
add_card(slide, 7.8, 4.35, 4.4, 0.9, "对比算法", "普通梯形 / 普通辛普森 / 复化梯形 / 复化辛普森", ORANGE)
add_card(slide, 7.8, 5.45, 4.4, 0.9, "展示风格", "多图 · 多表 · 少文字", PURPLE)

add_footer(slide, 1)


# Slide 2 工程背景
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "工程背景：位移 = 速度曲线下的面积",
          "非线性速度函数在 [0,1] 秒内的定积分")

add_picture(slide, velocity_img, 0.75, 1.35, width=7.2)

data = [
    ["项目", "数值"],
    ["v(t)", "4 / (1 + t²)"],
    ["积分区间", "[0, 1] s"],
    ["精确位移", "π m"],
    ["误差定义", "|近似值 - π|"],
]
add_table(slide, data, 8.35, 1.45, 3.9, 2.2, font_size=12, header_color=BLUE)

add_card(slide, 8.35, 4.05, 3.9, 0.85, "核心任务", "比较算法精度", ORANGE)
add_card(slide, 8.35, 5.05, 3.9, 0.85, "关键变量", "步长 h = 1/n", GREEN)
add_card(slide, 8.35, 6.05, 3.9, 0.85, "工程判断", "是否满足 10⁻³ m", PURPLE)

add_footer(slide, 2)


# Slide 3 算法示意
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "四种算法：从整段近似到分段近似",
          "直线近似与抛物线近似是精度差异的根源")

add_picture(slide, method_img, 0.65, 1.25, width=7.5)

method_table = [
    ["方法", "节点", "近似", "误差特征"],
    ["普通梯形", "2", "整段直线", "较大"],
    ["普通辛普森", "3", "整段抛物线", "较小"],
    ["复化梯形", "n+1", "分段直线", "O(h²)"],
    ["复化辛普森", "n+1", "分段抛物线", "一般 O(h⁴)"],
]
add_table(slide, method_table, 8.35, 1.45, 4.2, 3.0, font_size=11, header_color=BLUE)

add_card(slide, 8.35, 4.75, 4.2, 0.85, "本题特殊现象", "复化辛普森表现接近 O(h⁶)", PURPLE)
add_card(slide, 8.35, 5.75, 4.2, 0.85, "主要原因", "h⁴ 误差主项抵消", ORANGE)

add_footer(slide, 3)


# Slide 4 实验流程
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "实验流程", "从公式推导到工程精度判断")

steps = [
    ("建立模型\nS=∫v(t)dt", BLUE),
    ("编写算法\n4类公式", GREEN),
    ("改变 n\n计算误差", ORANGE),
    ("绘制图表\n观察规律", PURPLE),
    ("工程判定\n是否达标", RED),
]

x0 = 0.75
for i, (txt, color) in enumerate(steps):
    add_flow_box(slide, x0 + i * 2.45, 2.6, 1.8, 1.05, txt, color)
    if i < len(steps) - 1:
        add_arrow(slide, x0 + i * 2.45 + 1.85, 2.85, 0.45, 0.5)

add_card(slide, 1.1, 4.6, 3.2, 1.0, "输入", "v(t)=4/(1+t²)", BLUE)
add_card(slide, 5.0, 4.6, 3.2, 1.0, "输出", "数值结果 + 绝对误差", GREEN)
add_card(slide, 8.9, 4.6, 3.2, 1.0, "评价", "精度 / 收敛 / 计算量", ORANGE)

add_footer(slide, 4)


# Slide 5 普通公式结果
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "普通公式结果：辛普森明显优于梯形",
          "两种公式都只在整个区间上求积一次")

single_table = [
    ["方法", "数值结果", "绝对误差"],
    ["普通梯形", f"{T_single:.10f}", f"{err_T_single:.3e}"],
    ["普通辛普森", f"{S_single:.10f}", f"{err_S_single:.3e}"],
]
add_table(slide, single_table, 0.85, 1.35, 5.2, 1.55, font_size=13, header_color=BLUE)

add_picture(slide, bar_img, 6.35, 1.25, width=6.2)

add_card(slide, 0.95, 3.45, 4.95, 0.95, "结论 1", "普通辛普森误差约为普通梯形的 5.8%", GREEN)
add_card(slide, 0.95, 4.65, 4.95, 0.95, "结论 2", "抛物线近似更适合非线性函数", ORANGE)
add_card(slide, 0.95, 5.85, 4.95, 0.95, "但仍不足", "普通辛普森未达到 10⁻³ m", RED)

add_footer(slide, 5)


# Slide 6 复化公式数据表
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "复化求积结果总览", "等分数 n 增大，步长 h 减小，误差快速下降")

comp_table = [["n", "复化梯形 Tn", "误差T", "复化辛普森 Sn", "误差S"]]
for n, tv, te, sv, se in zip(ns, trap_values, trap_errors, simp_values, simp_errors):
    comp_table.append([
        str(n),
        f"{tv:.12f}",
        f"{te:.3e}",
        f"{sv:.12f}",
        f"{se:.3e}",
    ])

add_table(slide, comp_table, 0.55, 1.25, 12.2, 4.25, font_size=10, header_color=BLUE)

add_card(slide, 1.0, 5.95, 3.4, 0.85, "复化梯形", "n=16 达到 10⁻³ m", BLUE)
add_card(slide, 5.0, 5.95, 3.4, 0.85, "复化辛普森", "n=4 已达到 10⁻³ m", ORANGE)
add_card(slide, 9.0, 5.95, 3.4, 0.85, "综合表现", "辛普森精度最高", GREEN)

add_footer(slide, 6)


# Slide 7 误差—步长图
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "误差—步长关系", "步长 h 越小，复化公式误差越低")

add_picture(slide, error_step_img, 0.75, 1.15, width=8.6)

add_card(slide, 9.75, 1.55, 2.7, 0.9, "复化梯形", "≈ O(h²)", BLUE)
add_card(slide, 9.75, 2.75, 2.7, 0.9, "复化辛普森", "本题 ≈ O(h⁶)", ORANGE)
add_card(slide, 9.75, 3.95, 2.7, 0.9, "普通公式", "误差不随 n 变", RED)
add_card(slide, 9.75, 5.15, 2.7, 0.9, "图像结论", "辛普森下降最快", GREEN)

add_footer(slide, 7)


# Slide 8 误差—计算量图
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "误差—计算量对比", "相同计算量下，复化辛普森误差更小")

add_picture(slide, error_cost_img, 0.75, 1.15, width=8.6)

add_card(slide, 9.75, 1.55, 2.7, 0.9, "计算量", "约 O(n)", BLUE)
add_card(slide, 9.75, 2.75, 2.7, 0.9, "梯形", "稳定但收敛慢", ORANGE)
add_card(slide, 9.75, 3.95, 2.7, 0.9, "辛普森", "高精度高效率", GREEN)
add_card(slide, 9.75, 5.15, 2.7, 0.9, "推荐", "优先复化辛普森", PURPLE)

add_footer(slide, 8)


# Slide 9 收敛倍率
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "收敛速度验证", "n 翻倍时，误差缩小倍数反映误差阶")

add_picture(slide, ratio_img, 0.75, 1.15, width=7.5)

ratio_table = [["区间", "梯形误差比", "辛普森误差比"]]
for i in range(len(ns) - 1):
    ratio_table.append([
        f"{ns[i]}→{ns[i+1]}",
        f"{trap_errors[i] / trap_errors[i+1]:.1f}",
        f"{simp_errors[i] / simp_errors[i+1]:.1f}",
    ])

add_table(slide, ratio_table, 8.65, 1.35, 3.85, 3.55, font_size=10, header_color=BLUE)

add_card(slide, 8.65, 5.25, 3.85, 0.9, "理论观察", "梯形约 4 倍，辛普森后期约 64 倍", PURPLE)

add_footer(slide, 9)


# Slide 10 辛普森快速下降原因
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "为什么本题复化辛普森下降特别快？",
          "原因：四阶误差主项恰好抵消")

add_card(slide, 0.85, 1.25, 5.7, 1.15, "误差展开式",
         "Eₙ = -h⁴/180[f'''(1)-f'''(0)] + h⁶/1512[f⁽⁵⁾(1)-f⁽⁵⁾(0)] + O(h⁸)", BLUE)

derivative_table = [
    ["项目", "端点结果"],
    ["f'''(0)", "0"],
    ["f'''(1)", "0"],
    ["f⁽⁵⁾(0)", "0"],
    ["f⁽⁵⁾(1)", "60"],
]
add_table(slide, derivative_table, 7.2, 1.25, 4.3, 2.55, font_size=13, header_color=PURPLE)

add_card(slide, 0.85, 3.05, 5.7, 1.0, "关键抵消",
         "f'''(1)-f'''(0)=0 → h⁴ 项消失", RED)

add_card(slide, 0.85, 4.45, 5.7, 1.0, "真实主项",
         "Eₙ ≈ 5/126 · h⁶", GREEN)

add_card(slide, 7.2, 4.45, 4.3, 1.0, "步长减半",
         "误差约缩小 2⁶ = 64 倍", ORANGE)

add_footer(slide, 10)


# Slide 11 工程精度判定
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "工程精度判定", "以 10⁻³ m 和 10⁻⁴ m 为精度要求")

add_picture(slide, eng_img, 0.75, 1.15, width=7.7)

eng_table = [
    ["算法", "10⁻³ m", "10⁻⁴ m"],
    ["普通梯形", "未达标", "未达标"],
    ["普通辛普森", "未达标", "未达标"],
    ["复化梯形", "n=16", "n=64"],
    ["复化辛普森", "n=4", "n=4"],
]
add_table(slide, eng_table, 8.75, 1.45, 3.65, 2.8, font_size=12, header_color=BLUE)

add_card(slide, 8.75, 4.75, 3.65, 0.9, "工程推荐", "复化辛普森 n=4 即可", GREEN)
add_card(slide, 8.75, 5.85, 3.65, 0.9, "原因", "较少节点获得高精度", ORANGE)

add_footer(slide, 11)


# Slide 12 算法综合对比
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "四种算法综合对比", "精度、计算量、适用场景对比")

compare_table = [
    ["方法", "近似思想", "计算量", "精度表现", "适用场景"],
    ["普通梯形", "整段直线", "极低", "最低", "快速粗估"],
    ["普通辛普森", "整段抛物线", "很低", "中等", "初步精算"],
    ["复化梯形", "分段直线", "O(n)", "稳定收敛", "简单可靠"],
    ["复化辛普森", "分段抛物线", "O(n)", "最高", "工程优选"],
]
add_table(slide, compare_table, 0.65, 1.25, 12.1, 3.15, font_size=12, header_color=BLUE)

add_card(slide, 0.95, 5.05, 3.3, 0.95, "精度最高", "复化辛普森", GREEN)
add_card(slide, 5.0, 5.05, 3.3, 0.95, "实现最简单", "复化梯形", BLUE)
add_card(slide, 9.05, 5.05, 3.3, 0.95, "工程建议", "精度优先选辛普森", ORANGE)

add_footer(slide, 12)


# Slide 13 Python 程序结构
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "Python 程序结构", "模块化实现，便于复用和扩展")

code_table = [
    ["模块", "功能"],
    ["v(t)", "定义速度函数"],
    ["trapezoid_single()", "普通梯形公式"],
    ["simpson_single()", "普通辛普森公式"],
    ["composite_trapezoid()", "复化梯形公式"],
    ["composite_simpson()", "复化辛普森公式"],
    ["plot / ppt", "绘图并生成汇报 PPT"],
]
add_table(slide, code_table, 0.85, 1.25, 5.2, 4.3, font_size=12, header_color=BLUE)

add_flow_box(slide, 7.0, 1.55, 2.1, 0.8, "输入函数", BLUE)
add_arrow(slide, 9.25, 1.72, 0.55, 0.42)
add_flow_box(slide, 10.0, 1.55, 2.1, 0.8, "选择算法", GREEN)

add_flow_box(slide, 7.0, 3.0, 2.1, 0.8, "循环 n", ORANGE)
add_arrow(slide, 9.25, 3.17, 0.55, 0.42)
add_flow_box(slide, 10.0, 3.0, 2.1, 0.8, "计算误差", PURPLE)

add_flow_box(slide, 7.0, 4.45, 2.1, 0.8, "生成图表", RED)
add_arrow(slide, 9.25, 4.62, 0.55, 0.42)
add_flow_box(slide, 10.0, 4.45, 2.1, 0.8, "输出 PPT", BLUE)

add_footer(slide, 13)


# Slide 14 结论
slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "实验结论", "少文字总结版")

add_card(slide, 0.85, 1.35, 5.3, 1.0, "1. 普通公式", "辛普森明显优于梯形", BLUE)
add_card(slide, 7.0, 1.35, 5.3, 1.0, "2. 复化求积", "减小步长可显著降低误差", GREEN)

add_card(slide, 0.85, 2.9, 5.3, 1.0, "3. 收敛速度", "复化梯形约 O(h²)", ORANGE)
add_card(slide, 7.0, 2.9, 5.3, 1.0, "4. 本题特殊性", "复化辛普森表现接近 O(h⁶)", PURPLE)

add_card(slide, 0.85, 4.45, 5.3, 1.0, "5. 工程精度", "10⁻³ m 要求下，辛普森 n=4 即达标", RED)
add_card(slide, 7.0, 4.45, 5.3, 1.0, "6. 综合推荐", "精度优先选择复化辛普森", BLUE)

add_text_box(slide, 1.1, 6.35, 11.2, 0.45,
             "结论：复化辛普森公式在本实验中兼具高精度、快收敛和较低计算成本。",
             18, DARK, True, "center")

add_footer(slide, 14)


# 保存
prs.save(PPT_NAME)

print(f"已生成 PPT：{PPT_NAME}")
print(f"图片素材目录：{OUT_DIR}")