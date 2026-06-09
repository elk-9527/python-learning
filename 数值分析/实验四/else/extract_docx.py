import docx

doc = docx.Document(r'c:\Users\Lenovo\Desktop\报告.docx')
out = []
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        text = para.text.replace('\xa0', ' ')
        out.append(f'P{i}: [{para.style.name}] {text}')
out.append('---TABLES---')
for ti, table in enumerate(doc.tables):
    out.append(f'Table {ti}:')
    for ri, row in enumerate(table.rows):
        cells = [cell.text.replace('\xa0', ' ') for cell in row.cells]
        out.append(f'  Row{ri}: {cells}')

with open(r'c:\Users\Lenovo\Desktop\report_content.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('Done')
