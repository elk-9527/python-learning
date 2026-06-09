from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os


# =========================
# 基础设置
# =========================

PPT_NAME = "数值积分算法设计与实现PPT.pptx"
IMG_NAME = "误差步长与计算量对比图.png"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

TITLE_COLOR = RGBColor(31, 78, 121)
DARK_BLUE = RGBColor(31, 78, 121)
LIGHT_BLUE = RGBColor(221, 235, 247)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(90, 90, 90)
ORANGE = RGBColor(237, 125, 49)
GREEN = RGBColor(112, 173, 71)
RED = RGBColor(192, 0, 0)


def set_font(run, size=20, bold=False, color=BLACK):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_title_bar(slide, title, page=None):
    """添加顶部标题栏"""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        prs.slide_width,
        Inches(0.7)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK_BLUE
    bar.line.fill.background()

    textbox = slide.shapes.add_textbox(Inches(0.45), Inches(0.12), Inches(11.5), Inches(0.5))
    tf = textbox.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    set_font(run, size=24, bold=True, color=WHITE)

    if page is not None:
        footer = slide.shapes.add_textbox(Inches(12.2), Inches(0.18), Inches(0.9), Inches(0.4))
        tf2 = footer.text_frame
        tf2.clear()
        p2 = tf2.paragraphs[0]
        p2.text = str(page)
        p2.alignment = PP_ALIGN.RIGHT
        run2 = p2.runs[0]
        set_font(run2, size=14, color=WHITE)


def add_footer(slide):
    """添加页脚"""
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(7.25),
        prs.slide_width,
        Inches(0.03)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_BLUE
    line.line.fill.background()

    footer = slide.shapes.add_textbox(Inches(0.45), Inches(7.28), Inches(12.5), Inches(0.2))
    tf = footer.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "数值积分算法设计与实现 | 梯形公式、辛普森公式、复化求积"
    p.alignment = PP_ALIGN.RIGHT
    run = p.runs[0]
    set_font(run, size=9, color=GRAY)


def add_bullets(slide, left, top, width, height, bullets, font_size=20):
    """添加项目符号文本"""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()

    for i, item in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        p.text = item
        p.level = 0
        p.space_after = Pt(8)
        run = p.runs[0]
        set_font(run, size=font_size, color=BLACK)

    return box


def add_formula_box(slide, left, top, width, height, text, font_size=22, fill_color=LIGHT_BLUE):
    """添加公式框"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = DARK_BLUE
    shape.line.width = Pt(1.2)

    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    set_font(run, size=font_size, bold=True, color=DARK_BLUE)

    return shape


def add_section_tag(slide, left, top, text, color=ORANGE):
    """添加小标签"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(2.3),
        Inches(0.42)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    set_font(run, size=14, bold=True, color=WHITE)


def create_slide(title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, title, len(prs.slides))
    add_footer(slide)
    return slide


# =========================
# 第 1 页：封面
# =========================

slide = prs.slides.add_slide(prs.slide_layouts[6])

bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(242, 247, 252)
bg.line.fill.background()

title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.45), Inches(11.5), Inches(1.0))
tf = title_box.text_frame
tf.clear()
p = tf.paragraphs[0]
p.text = "数值积分算法设计与实现"
p.alignment = PP_ALIGN.CENTER
run = p.runs[0]
set_font(run, size=40, bold=True, color=DARK_BLUE)

sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.55), Inches(11.5), Inches(0.7))
tf = sub_box.text_frame
tf.clear()
p = tf.paragraphs[0]
p.text = "梯形公式、辛普森公式与复化求积"
p.alignment = PP_ALIGN.CENTER
run = p.runs[0]
set_font(run, size=24, bold=True, color=ORANGE)

formula = "S = ∫₀¹ 4/(1+t²) dt = π"
add_formula_box(slide, 3.1, 3.55, 7.2, 0.85, formula, font_size=26)

info = slide.shapes.add_textbox(Inches(1.0), Inches(5.3), Inches(11.5), Inches(0.8))
tf = info.text_frame
tf.clear()
p = tf.paragraphs[0]
p.text = "实验汇报 PPT"
p.alignment = PP_ALIGN.CENTER
run = p.runs[0]
set_font(run, size=20, color=GRAY)


# =========================
# 第 2 页：实验背景与目标
# =========================

slide = create_slide("一、实验背景与目标")

add_section_tag(slide, 0.7, 1.05, "工程背景")

add_bullets(
    slide,
    0.8, 1.65, 5.6, 2.4,
    [
        "某滑块做非线性变速直线运动",
        "瞬时速度函数：v(t)=4/(1+t²) m/s",
        "要求计算 0 到 1 秒内的位移",
        "位移由速度函数定积分给出"
    ],
    font_size=20
)

add_formula_box(
    slide,
    6.8, 1.55, 5.4, 1.0,
    "S = ∫₀¹ 4/(1+t²) dt",
    font_size=24
)

add_section_tag(slide, 0.7, 4.25, "实验目标", color=GREEN)

add_bullets(
    slide,
    0.8, 4.8, 11.5, 1.8,
    [
        "推导并实现普通梯形、普通辛普森、复化梯形、复化辛普森四种算法",
        "固定积分区间，改变等分份数 n，观察误差随步长 h 的变化",
        "对比四种算法的精度、收敛速度和计算量，并分析工程适用性"
    ],
    font_size=20
)


# =========================
# 第 3 页：精确值与误差标准
# =========================

slide = create_slide("二、精确值与误差标准")

add_formula_box(
    slide,
    1.0, 1.25, 11.3, 0.8,
    "∫ 4/(1+t²) dt = 4 arctan(t) + C",
    font_size=24
)

add_formula_box(
    slide,
    1.0, 2.45, 11.3, 0.8,
    "S = 4[arctan(1) - arctan(0)] = 4 × π/4 = π",
    font_size=24
)

add_bullets(
    slide,
    1.1, 3.7, 11.2, 1.8,
    [
        "精确位移：S = π ≈ 3.1415926536 m",
        "所有数值结果均与 π 比较",
        "绝对误差定义：e = |数值近似值 - π|"
    ],
    font_size=22
)

add_formula_box(
    slide,
    3.3, 5.85, 6.8, 0.65,
    "e = |Q - π|",
    font_size=26,
    fill_color=RGBColor(255, 242, 204)
)


# =========================
# 第 4 页：四种算法总览
# =========================

slide = create_slide("三、四种数值积分算法总览")

rows, cols = 5, 5
table_shape = slide.shapes.add_table(rows, cols, Inches(0.55), Inches(1.25), Inches(12.2), Inches(4.7))
table = table_shape.table

headers = ["方法", "近似思想", "节点数", "误差阶", "特点"]
data = [
    ["普通梯形", "直线近似整段曲线", "2", "O(h³)", "简单但误差较大"],
    ["普通辛普森", "抛物线近似整段曲线", "3", "O(h⁵)", "精度明显提高"],
    ["复化梯形", "分段直线近似", "n+1", "O(h²)", "稳定收敛"],
    ["复化辛普森", "分段抛物线近似", "n+1", "一般 O(h⁴)", "精度最高"]
]

for j, h in enumerate(headers):
    cell = table.cell(0, j)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = DARK_BLUE
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            set_font(r, size=14, bold=True, color=WHITE)

for i, row in enumerate(data, start=1):
    for j, value in enumerate(row):
        cell = table.cell(i, j)
        cell.text = value
        if i % 2 == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(235, 242, 250)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_font(r, size=13, color=BLACK)

add_bullets(
    slide,
    0.8, 6.25, 12.0, 0.6,
    [
        "核心规律：复化求积通过减小步长 h，提高对曲线的局部近似精度。"
    ],
    font_size=20
)


# =========================
# 第 5 页：普通梯形与普通辛普森
# =========================

slide = create_slide("四、普通梯形公式与普通辛普森公式")

add_section_tag(slide, 0.7, 1.05, "普通梯形")

add_formula_box(
    slide,
    0.8, 1.65, 5.6, 0.75,
    "T = (b-a)/2 · [f(a)+f(b)]",
    font_size=20
)

add_bullets(
    slide,
    0.9, 2.65, 5.3, 1.5,
    [
        "本题中 f(0)=4，f(1)=2",
        "T = 1/2 × (4+2) = 3",
        "误差 ≈ 0.1415926536"
    ],
    font_size=18
)

add_section_tag(slide, 6.9, 1.05, "普通辛普森", color=GREEN)

add_formula_box(
    slide,
    6.9, 1.65, 5.7, 0.75,
    "S = (b-a)/6 · [f(a)+4f(m)+f(b)]",
    font_size=18
)

add_bullets(
    slide,
    7.0, 2.65, 5.3, 1.5,
    [
        "中点 m=0.5，f(0.5)=3.2",
        "S = 3.1333333333",
        "误差 ≈ 0.0082593203"
    ],
    font_size=18
)

add_formula_box(
    slide,
    2.0, 5.3, 9.3, 0.75,
    "结论：辛普森公式用抛物线近似，精度明显优于梯形直线近似",
    font_size=20,
    fill_color=RGBColor(226, 239, 218)
)


# =========================
# 第 6 页：复化梯形公式
# =========================

slide = create_slide("五、复化梯形公式")

add_formula_box(
    slide,
    0.9, 1.15, 11.5, 0.75,
    "h = (b-a)/n，xᵢ = a + ih",
    font_size=23
)

add_formula_box(
    slide,
    0.9, 2.25, 11.5, 0.85,
    "Tₙ = h[(f(x₀)+f(xₙ))/2 + Σ f(xᵢ)]",
    font_size=23,
    fill_color=RGBColor(255, 242, 204)
)

add_bullets(
    slide,
    1.0, 3.55, 11.4, 2.2,
    [
        "将区间 [0,1] 等分为 n 个小区间",
        "每个小区间用梯形公式近似",
        "全局误差阶：E_Tₙ = O(h²)",
        "当 n 翻倍时，h 减半，误差约缩小为原来的 1/4"
    ],
    font_size=21
)


# =========================
# 第 7 页：复化辛普森公式
# =========================

slide = create_slide("六、复化辛普森公式")

add_formula_box(
    slide,
    0.9, 1.1, 11.5, 0.75,
    "n 必须为偶数，每两个小区间组成一个辛普森面板",
    font_size=21
)

add_formula_box(
    slide,
    0.7, 2.15, 12.0, 1.05,
    "Sₙ = h/3 [f(x₀)+f(xₙ)+4Σf(xᵢ奇)+2Σf(xᵢ偶)]",
    font_size=20,
    fill_color=RGBColor(255, 242, 204)
)

add_bullets(
    slide,
    1.0, 3.65, 11.4, 2.1,
    [
        "复化辛普森公式使用分段抛物线逼近原函数",
        "一般全局误差阶：E_Sₙ = O(h⁴)",
        "当 n 翻倍时，一般误差约缩小为原来的 1/16",
        "本题中因误差主项抵消，实际表现出接近 O(h⁶) 的收敛特征"
    ],
    font_size=21
)


# =========================
# 第 8 页：程序实现
# =========================

slide = create_slide("七、Python 程序实现")

add_bullets(
    slide,
    0.8, 1.15, 5.8, 4.7,
    [
        "定义速度函数 v(t)=4/(1+t²)",
        "使用 math.pi 作为精确值",
        "分别编写四个独立函数：",
        "trapezoid_single()",
        "simpson_single()",
        "composite_trapezoid()",
        "composite_simpson()"
    ],
    font_size=19
)

code_box = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(6.8),
    Inches(1.25),
    Inches(5.8),
    Inches(4.7)
)
code_box.fill.solid()
code_box.fill.fore_color.rgb = RGBColor(245, 245, 245)
code_box.line.color.rgb = RGBColor(180, 180, 180)

code = (
    "def composite_simpson(f,a,b,n):\n"
    "    if n % 2 != 0:\n"
    "        raise ValueError('n必须为偶数')\n\n"
    "    h = (b-a)/n\n"
    "    odd_sum = 0\n"
    "    even_sum = 0\n\n"
    "    for i in range(1,n):\n"
    "        x = a + i*h\n"
    "        if i % 2 == 1:\n"
    "            odd_sum += f(x)\n"
    "        else:\n"
    "            even_sum += f(x)\n\n"
    "    return h/3*(f(a)+f(b)\n"
    "        +4*odd_sum+2*even_sum)"
)

tf = code_box.text_frame
tf.clear()
p = tf.paragraphs[0]
p.text = code
p.alignment = PP_ALIGN.LEFT
run = p.runs[0]
run.font.name = "Consolas"
run.font.size = Pt(12)
run.font.color.rgb = BLACK


# =========================
# 第 9 页：实验数据
# =========================

slide = create_slide("八、实验结果数据")

rows, cols = 7, 5
table_shape = slide.shapes.add_table(rows, cols, Inches(0.4), Inches(1.1), Inches(12.5), Inches(5.2))
table = table_shape.table

headers = ["n", "复化梯形值", "梯形误差", "复化辛普森值", "辛普森误差"]
data = [
    ["2", "3.100000000000", "4.159265e-02", "3.133333333333", "8.259320e-03"],
    ["4", "3.131176470588", "1.041618e-02", "3.141568627451", "2.402614e-05"],
    ["8", "3.138988494491", "2.604159e-03", "3.141592502459", "1.511311e-07"],
    ["16", "3.140941612041", "6.510415e-04", "3.141592651225", "2.364971e-09"],
    ["32", "3.141429893174", "1.627604e-04", "3.141592653553", "3.70e-11"],
    ["64", "3.141551963485", "4.069010e-05", "3.141592653589", "5.8e-13"],
]

for j, h in enumerate(headers):
    cell = table.cell(0, j)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = DARK_BLUE
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            set_font(r, size=12, bold=True, color=WHITE)

for i, row in enumerate(data, start=1):
    for j, value in enumerate(row):
        cell = table.cell(i, j)
        cell.text = value
        if i % 2 == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(235, 242, 250)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_font(r, size=11, color=BLACK)

add_bullets(
    slide,
    0.7, 6.55, 12.0, 0.5,
    [
        "数据表明：n 增大时，复化梯形和复化辛普森误差均下降，其中复化辛普森下降更快。"
    ],
    font_size=18
)


# =========================
# 第 10 页：误差图
# =========================

slide = create_slide("九、误差—步长与计算量对比图")

img_path = IMG_NAME

if os.path.exists(img_path):
    slide.shapes.add_picture(img_path, Inches(0.45), Inches(1.05), width=Inches(12.4))
else:
    warning = slide.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(10.8), Inches(1.0))
    tf = warning.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = f"未找到图片：{IMG_NAME}\n请将图片与 make_ppt.py 放在同一文件夹中。"
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    set_font(run, size=24, bold=True, color=RED)


# =========================
# 第 11 页：图像分析
# =========================

slide = create_slide("十、图像结果分析")

add_section_tag(slide, 0.7, 1.05, "误差—步长图")

add_bullets(
    slide,
    0.8, 1.6, 12.0, 2.15,
    [
        "随着 n 增大，步长 h=1/n 减小，复化公式误差明显下降",
        "普通梯形与普通辛普森误差不随 n 变化，在图中表现为水平线",
        "复化梯形误差曲线接近 O(h²) 参考线，符合二阶收敛规律",
        "复化辛普森误差下降更快，本题中接近 O(h⁶) 参考线"
    ],
    font_size=19
)

add_section_tag(slide, 0.7, 4.1, "误差—计算量图", color=GREEN)

add_bullets(
    slide,
    0.8, 4.65, 12.0, 1.65,
    [
        "复化梯形和复化辛普森的计算量均约为 O(n)",
        "在相同或相近节点数下，复化辛普森误差远小于复化梯形",
        "复化辛普森公式具有更高的精度和工程计算效率"
    ],
    font_size=19
)


# =========================
# 第 12 页：本题特殊误差现象
# =========================

slide = create_slide("十一、本题复化辛普森误差下降更快的原因")

add_formula_box(
    slide,
    0.8, 1.15, 11.8, 0.75,
    "一般复化辛普森误差：Eₙ = O(h⁴)",
    font_size=24
)

add_formula_box(
    slide,
    0.8, 2.2, 11.8, 1.0,
    "Eₙ = -h⁴/180 [f'''(1)-f'''(0)] + h⁶/1512 [f⁽⁵⁾(1)-f⁽⁵⁾(0)] + O(h⁸)",
    font_size=17,
    fill_color=RGBColor(255, 242, 204)
)

add_bullets(
    slide,
    1.0, 3.6, 11.4, 2.2,
    [
        "本题 f(x)=4/(1+x²)",
        "三阶导数 f'''(x)=96x(1-x²)/(1+x²)⁴",
        "端点满足：f'''(0)=0，f'''(1)=0",
        "因此 h⁴ 主误差项消失，主项变为 h⁶"
    ],
    font_size=20
)

add_formula_box(
    slide,
    3.1, 6.1, 7.1, 0.7,
    "Eₙ = 5/126 · h⁶ + O(h⁸)",
    font_size=24,
    fill_color=RGBColor(226, 239, 218)
)


# =========================
# 第 13 页：工程精度分析
# =========================

slide = create_slide("十二、工程精度分析")

add_formula_box(
    slide,
    1.0, 1.05, 11.3, 0.7,
    "设基础工程精度要求：绝对误差 < 10⁻³ m",
    font_size=24,
    fill_color=RGBColor(255, 242, 204)
)

rows, cols = 5, 3
table_shape = slide.shapes.add_table(rows, cols, Inches(1.1), Inches(2.1), Inches(11.1), Inches(3.0))
table = table_shape.table

headers = ["方法", "误差表现", "是否满足 10⁻³ m"]
data = [
    ["普通梯形", "1.42×10⁻¹", "不满足"],
    ["普通辛普森", "8.26×10⁻³", "不满足"],
    ["复化梯形", "n=16 时 6.51×10⁻⁴", "满足"],
    ["复化辛普森", "n=4 时 2.40×10⁻⁵", "满足"],
]

for j, h in enumerate(headers):
    cell = table.cell(0, j)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = DARK_BLUE
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            set_font(r, size=13, bold=True, color=WHITE)

for i, row in enumerate(data, start=1):
    for j, value in enumerate(row):
        cell = table.cell(i, j)
        cell.text = value
        if "不满足" in value:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 230, 230)
        elif "满足" in value:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(226, 239, 218)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_font(r, size=13, color=BLACK)

add_bullets(
    slide,
    1.0, 5.65, 11.4, 0.9,
    [
        "复化辛普森公式用较少节点即可达到较高精度，工程应用价值最高。"
    ],
    font_size=21
)


# =========================
# 第 14 页：总结
# =========================

slide = create_slide("十三、实验总结")

add_bullets(
    slide,
    0.9, 1.15, 11.8, 4.8,
    [
        "普通辛普森公式精度明显高于普通梯形公式",
        "复化求积通过减小步长显著降低截断误差",
        "复化梯形公式为二阶方法，误差满足 O(h²)",
        "复化辛普森公式一般为四阶方法，误差满足 O(h⁴)",
        "本题因 f'''(0)=f'''(1)=0，复化辛普森实际表现为六阶收敛",
        "综合精度、收敛速度和计算量，复化辛普森公式表现最好"
    ],
    font_size=22
)

add_formula_box(
    slide,
    2.0, 6.2, 9.3, 0.75,
    "最终结论：复化辛普森公式是本题最优的数值积分方法",
    font_size=23,
    fill_color=RGBColor(226, 239, 218)
)


# =========================
# 保存 PPT
# =========================

prs.save(PPT_NAME)

print(f"已生成 PPT 文件：{PPT_NAME}")